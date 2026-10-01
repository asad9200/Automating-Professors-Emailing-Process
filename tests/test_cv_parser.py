from io import BytesIO

import pytest
from docx import Document

from profreach.tools import cv_parser
from profreach.tools.cv_parser import CVParseError, extract_cv_text


def test_extracts_text_from_docx_paragraphs_and_tables():
    document = Document()
    document.add_paragraph("Research in machine learning")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "University of Example"
    buffer = BytesIO()
    document.save(buffer)

    text = extract_cv_text(buffer.getvalue(), "cv.DOCX")

    assert "Research in machine learning" in text
    assert "University of Example" in text


def test_extracts_text_from_pdf(monkeypatch):
    class Page:
        def extract_text(self):
            return "Computer vision publications"

    class Reader:
        def __init__(self, stream):
            assert stream.read() == b"pdf bytes"
            self.pages = [Page()]

    monkeypatch.setattr(cv_parser, "PdfReader", Reader)

    assert extract_cv_text(b"pdf bytes", "cv.pdf") == "Computer vision publications"


@pytest.mark.parametrize(
    ("file_bytes", "filename", "message"),
    [
        (b"", "cv.pdf", "empty"),
        (b"content", "cv.txt", "PDF or DOCX"),
    ],
)
def test_rejects_empty_or_unsupported_files(file_bytes, filename, message):
    with pytest.raises(CVParseError, match=message):
        extract_cv_text(file_bytes, filename)


def test_rejects_pdf_without_extractable_text(monkeypatch):
    class EmptyPage:
        def extract_text(self):
            return None

    class Reader:
        def __init__(self, stream):
            self.pages = [EmptyPage()]

    monkeypatch.setattr(cv_parser, "PdfReader", Reader)

    with pytest.raises(CVParseError, match="No text was found"):
        extract_cv_text(b"pdf bytes", "scan.pdf")
