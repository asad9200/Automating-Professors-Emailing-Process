from io import BytesIO

from docx import Document
from pypdf import PdfReader

MAX_CV_SIZE_BYTES = 10 * 1024 * 1024


class CVParseError(ValueError):
    """Raised when an uploaded CV is unsupported or cannot be read."""


def extract_cv_text(file_bytes: bytes, filename: str) -> str:
    if not file_bytes:
        raise CVParseError("The uploaded file is empty.")
    if len(file_bytes) > MAX_CV_SIZE_BYTES:
        raise CVParseError("The CV must be smaller than 10 MB.")

    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    try:
        if extension == "pdf":
            pages = PdfReader(BytesIO(file_bytes)).pages
            text = "\n".join(page.extract_text() or "" for page in pages)
        elif extension == "docx":
            document = Document(BytesIO(file_bytes))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            table_cells = [
                cell.text for table in document.tables for row in table.rows for cell in row.cells
            ]
            text = "\n".join(paragraphs + table_cells)
        else:
            raise CVParseError("Upload a PDF or DOCX file.")
    except CVParseError:
        raise
    except Exception as error:
        raise CVParseError("The uploaded CV could not be read.") from error

    text = text.strip()
    if not text:
        raise CVParseError("No text was found. Scanned PDFs need OCR before they can be used.")
    return text
